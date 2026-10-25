"""Release-scoped inventory checks and publish-after-success document refresh."""
from __future__ import annotations

from concurrent.futures import ThreadPoolExecutor
from copy import deepcopy
from datetime import datetime, timezone
import base64
import hashlib
import io
import json
import os
from pathlib import Path, PurePosixPath
import re
import shutil
import subprocess
import sys
import tempfile
from uuid import uuid4
import zipfile

from docx import Document
from docx.oxml.ns import qn

from spec_retriever.spec_store.manifest import DocumentSpec, Manifest, load_manifest, sha256_file
from spec_retriever.sync.official import OfficialCatalog, fetch_bytes
from spec_retriever.retrieval.references import bibliography


def reference_inventory(corpus, baselines: list[str]) -> list[dict]:
    """Discover cited specification numbers; this is one hop, not applicability."""
    result = []
    for baseline in dict.fromkeys(baselines):
        for document in corpus.status(baseline)["documents"]:
            if document["status"] != "indexed":
                continue
            records = list(corpus.clause_records(baseline, document["spec"], document["version"]))
            found = bibliography(records)
            known = {value.get("spec") for value in found.values()}
            for number, value in found.items():
                if value.get("spec"):
                    # Preserve multipart names instead of silently turning
                    # TS 38.101-1 into the different identifier TS 38.101.
                    names = set(re.findall(r"(?:TS|TR)\s+(\d{2,3}\.\d{3}(?:-\d+)?)\b", value.get('raw', ''), re.I))
                    for spec in sorted(names or {value['spec']}):
                        result.append({"source_baseline": baseline, "source_spec": document["spec"],
                                       "reference": number, "spec": spec,
                                       **{k: value.get(k) for k in ("version", "release")}})
                        known.add(spec)
            # Explicit mentions outside the bibliography are also inventory leads.
            for record in records:
                for block in record["blocks"]:
                    for spec in re.findall(r"(?:3GPP\s+)?(?:TS|TR)\s+(\d{2,3}\.\d{3}(?:-\d+)?)\b", block["text"], re.I):
                        if spec not in known and spec != document["spec"]:
                            known.add(spec)
                            result.append({"source_baseline": baseline, "source_spec": document["spec"],
                                           "source_clause": record["clause_id"], "spec": spec,
                                           "version": None, "release": None})
    return result


def missing_references(references, available, release):
    gaps = []
    for ref in references:
        candidates = [doc for doc in available if doc['spec'] == ref['spec'] and doc['status'] == 'indexed']
        if ref.get('version'):
            candidates = [doc for doc in candidates if doc['version'] == ref['version']]
        else:
            wanted = ref.get('release') or release
            candidates = [doc for doc in candidates if int(doc['version'].split('.')[0]) == wanted]
        if not candidates:
            gaps.append({**ref, 'reason': 'exact_version_missing' if ref.get('version') else 'document_missing'})
    return gaps


def check_documents(corpus, manifest: Manifest, *, catalog=None, required_specs=(),
                    dependencies=(), include_references=False, series=()) -> dict:
    catalog = catalog or OfficialCatalog()
    scopes = list(dict.fromkeys([manifest.baseline, *dependencies]))
    try:
        state = corpus.status(manifest.baseline)
    except ValueError:
        state = {"documents": [], "manifest_sha256": None}
        scopes.remove(manifest.baseline)
    references = reference_inventory(corpus, scopes)
    available = [doc for scope in scopes for doc in corpus.status(scope)["documents"]]
    gaps = missing_references(references, available, manifest.release)
    selected = {item.spec for item in manifest.documents} | set(required_specs)
    if include_references:
        selected.update(ref["spec"] for ref in references if not ref.get("version")
                        and ref.get("release") in (None, manifest.release))
    # A series comparison includes only specifications that publish this Release;
    # other series entries remain visible as not_available_for_release.
    for number in series:
        selected.update(catalog.series_specs(number))
    local = {item.spec: item for item in manifest.documents}
    indexed = {item["spec"]: item for item in state["documents"]}

    def inspect(spec):
        item = local.get(spec)
        row = {"spec": spec, "local_version": item.version if item else None,
               "local_file_status": "not_in_manifest", "index_status": indexed.get(spec, {}).get("status", "not_indexed")}
        if item:
            path = manifest.document_path(item)
            row["local_file_status"] = ("missing" if not path.is_file() else
                                        "ok" if sha256_file(path) == item.sha256 else "checksum_mismatch")
        try:
            latest = catalog.latest(spec, manifest.release)
            row.update(latest_version=latest.version if latest else None,
                       source_url=latest.url if latest else None,
                       version_status=("not_available_for_release" if not latest else "missing" if not item
                                       else "current" if item.version == latest.version else
                                       "outdated" if tuple(map(int, item.version.split('.'))) < tuple(map(int, latest.version.split('.')))
                                       else "local_ahead_of_catalog"))
        except (ValueError, OSError) as error:
            row.update(version_status="check_failed", error=str(error))
        return row

    with ThreadPoolExecutor(max_workers=6) as executor:
        documents = list(executor.map(inspect, sorted(selected)))
    # A successful directory response does not verify upstream archive bytes.
    mandatory = {item.spec for item in manifest.documents} | set(required_specs)
    applicable = [row for row in documents if row['version_status'] != 'not_available_for_release' or row['spec'] in mandatory]
    ready = bool(applicable) and all(row["version_status"] == "current" and row["local_file_status"] == "ok"
                and row["index_status"] == "indexed" for row in applicable)
    return {"baseline": manifest.baseline, "release": manifest.release,
            "checked_at": catalog.checked_at, "version_policy": "latest_within_same_release",
            "archive_bytes_revalidated": False, "manifest_matches_index": state["manifest_sha256"] == manifest.digest,
            "documents": documents, "missing_references": gaps,
            "all_selected_current": ready and state["manifest_sha256"] == manifest.digest,
            "reference_scan_complete": all(doc["status"] == "indexed" for doc in available) and bool(state["documents"])
                                       and state["manifest_sha256"] == manifest.digest,
            "reference_depth": 1, "reference_inventory": references,
            "coverage_scope": {"required_specs": sorted(set(required_specs)), "series": list(series),
                               "include_references": include_references, "global_3gpp_completeness": "not_assessed"}}


def extract_document(archive: bytes | Path, spec: str, destination: Path) -> tuple[Path, dict]:
    """Extract a single document or merge an official section-split bundle."""
    try:
        with zipfile.ZipFile(archive if isinstance(archive, Path) else io.BytesIO(archive)) as bundle:
            prefix = spec.replace('.', '') + '-'
            candidates = [item for item in bundle.infolist()
                          if PurePosixPath(item.filename).name.lower().startswith(prefix)
                          and PurePosixPath(item.filename).suffix.lower() in {".doc", ".docx"}
                          and not item.is_dir()]
            if not candidates:
                raise ValueError(f"{spec}: official ZIP contains no specification DOC/DOCX")
            candidates.sort(key=lambda item: PurePosixPath(item.filename).name.lower())
            if sum(item.file_size for item in candidates) > 512 * 1024 * 1024:
                raise ValueError("uncompressed specification documents exceed 512 MiB")
            members = []
            paths = []
            for index, entry in enumerate(candidates):
                suffix = PurePosixPath(entry.filename).suffix.lower()
                target = destination / (("source" if len(candidates) == 1 else f"part-{index:03d}") + suffix)
                target.unlink(missing_ok=True)
                data = bundle.read(entry)
                with target.open("xb") as stream:
                    stream.write(data)
                members.append({"name": entry.filename, "sha256": hashlib.sha256(data).hexdigest()})
                paths.append(target)
    except zipfile.BadZipFile as error:
        raise ValueError(f"{spec}: invalid official ZIP") from error
    archive_hash = sha256_file(archive) if isinstance(archive, Path) else hashlib.sha256(archive).hexdigest()
    if len(paths) == 1:
        return paths[0], {"archive_sha256": archive_hash,
                          "archive_member": members[0]["name"], "original_sha256": members[0]["sha256"]}
    converted = [convert_legacy_doc(path) if path.suffix == ".doc" else path for path in paths]
    merged = destination / "source.docx"
    merged.unlink(missing_ok=True)
    ranges = merge_docx_parts(converted, merged)
    for part, span in zip(members, ranges):
        part["body_start"], part["body_end"] = span
    return merged, {"archive_sha256": archive_hash, "archive_members": members,
                    "conversion": "merged_official_parts"}


def merge_docx_parts(parts: list[Path], output: Path) -> list[tuple[int, int]]:
    """Merge ordered source bodies; verify that no supported text block is lost."""
    if len(parts) < 2:
        raise ValueError("multi-part merge requires at least two documents")
    main = Document(parts[0])
    bodies = [main.element.body]
    source_text = []

    def body_elements(body):
        return [element for element in body if element.tag != qn("w:sectPr")]

    def texts(elements):
        return ["".join(element.xpath(".//w:t/text()")) for element in elements
                if element.tag in {qn("w:p"), qn("w:tbl")}]

    first = body_elements(bodies[0])
    source_text.extend(texts(first))
    ranges = [(0, len(first))]
    position = len(first)
    for path in parts[1:]:
        body = Document(path).element.body
        elements = body_elements(body)
        source_text.extend(texts(elements))
        start = position
        for element in elements:
            main.element.body.insert(position, deepcopy(element))
            position += 1
        ranges.append((start, position))
    if texts(body_elements(main.element.body)) != source_text:
        raise ValueError("multi-part merge changed source paragraph or table text")
    main.save(output)
    if texts(body_elements(Document(output).element.body)) != source_text:
        output.unlink(missing_ok=True)
        raise ValueError("saved multi-part DOCX changed source text")
    return ranges


def convert_legacy_doc(source: Path) -> Path:
    """Convert an official legacy DOC with LibreOffice or Windows Word."""
    output = source.with_suffix('.docx')
    if output.exists():
        raise ValueError("conversion destination already exists")
    libreoffice = shutil.which("libreoffice") or shutil.which("soffice")
    if libreoffice:
        with tempfile.TemporaryDirectory() as profile:
            try:
                completed = subprocess.run([libreoffice, "-env:UserInstallation=" + Path(profile).as_uri(),
                    "--headless", "--convert-to", "docx", "--outdir", str(source.parent), str(source)],
                    capture_output=True, timeout=180)
            except subprocess.TimeoutExpired as error:
                raise ValueError("legacy DOC conversion timed out") from error
        if completed.returncode or not output.is_file():
            raise ValueError("legacy DOC conversion failed; active corpus was not changed")
        return output
    if sys.platform != "win32":
        raise ValueError("legacy DOC requires LibreOffice; active corpus was not changed")
    quote = lambda path: "'" + str(path.resolve()).replace("'", "''") + "'"
    script = f"""$ErrorActionPreference='Stop'; $word=$null; $doc=$null
try {{
 $word=New-Object -ComObject Word.Application
 $word.Visible=$false; $word.DisplayAlerts=0; $word.AutomationSecurity=3
 $doc=$word.Documents.Open({quote(source)},$false,$true,$false)
 $doc.SaveAs2({quote(output)},16)
}} finally {{
 if($null -ne $doc) {{ $doc.Close(0); [void][Runtime.InteropServices.Marshal]::FinalReleaseComObject($doc) }}
 if($null -ne $word) {{ $word.Quit(); [void][Runtime.InteropServices.Marshal]::FinalReleaseComObject($word) }}
}}
"""
    encoded = base64.b64encode(script.encode('utf-16le')).decode('ascii')
    try:
        completed = subprocess.run(['powershell', '-NoProfile', '-NonInteractive', '-EncodedCommand', encoded],
                                   capture_output=True, timeout=180)
    except subprocess.TimeoutExpired as error:
        raise ValueError("legacy DOC conversion timed out") from error
    if completed.returncode or not output.is_file():
        raise ValueError("legacy DOC conversion failed; active corpus was not changed")
    return output


def sync_documents(corpus, manifest: Manifest, directory: str | Path, *, catalog=None,
                   required_specs=(), dependencies=(), include_references=False, series=(),
                   embed=None, lexical_only=False, fetch=fetch_bytes, converter=convert_legacy_doc,
                   embedding_profile=None, force=False, embedding_current=None) -> dict:
    """Publish a fresh generation only after all selected documents and models succeed."""
    if not lexical_only and embed is None:
        raise ValueError("embedding callback required, or explicitly select lexical_only")
    root = Path(directory).resolve()
    root.mkdir(parents=True, exist_ok=True)
    # Serialise publishers across processes. A crash leaves a visible lock that
    # requires inspection; never guess that another process has stopped.
    lock = root / '.sync.lock'
    with lock.open('x', encoding='utf-8') as stream:
        stream.write(str(os.getpid()))
    try:
        catalog = catalog or OfficialCatalog()
        check = check_documents(corpus, manifest, catalog=catalog, required_specs=required_specs,
                                dependencies=dependencies, include_references=include_references, series=series)
        mandatory = {item.spec for item in manifest.documents} | set(required_specs)
        failures = [row for row in check['documents'] if row['version_status'] in
                    {'check_failed', 'local_ahead_of_catalog'} or
                    (row['version_status'] == 'not_available_for_release' and row['spec'] in mandatory)]
        if failures:
            raise ValueError("official version check failed; active corpus was not changed: " +
                             ', '.join(row['spec'] for row in failures))
        active_path = root / 'current.json'
        if not force and check['all_selected_current'] and active_path.is_file():
            active = json.loads(active_path.read_text(encoding='utf-8'))
            if (active.get('baseline') == manifest.baseline and active.get('manifest_sha256') == manifest.digest
                    and active.get('embedding_profile') == embedding_profile
                    and (lexical_only or (active.get('embedding', {}).get('status') == 'complete'
                         and (embedding_current is None or embedding_current(corpus, manifest.baseline))))):
                return {**active, 'checked_at': catalog.checked_at, 'published': False, 'unchanged': True,
                        'missing_references': check['missing_references']}
        selected = [row for row in check['documents'] if row['latest_version']]
        if not selected:
            raise ValueError("no documents available for the selected Release")
        baseline = f"sync-r{manifest.release}-" + datetime.now(timezone.utc).strftime('%Y%m%dT%H%M%S') + '-' + uuid4().hex[:8]
        generation = root / baseline
        generation.mkdir()
        items, sources = [], []
        previous = {item.spec: item for item in manifest.documents}
        for row in selected:
            spec = row['spec']
            folder = generation / spec
            folder.mkdir()
            old = previous.get(spec)
            # Healthy unchanged DOCX can be reused; changed or missing bytes are
            # downloaded anew. Every generation is re-ingested and re-embedded.
            if old and row['version_status'] == 'current' and row['local_file_status'] == 'ok':
                data = manifest.document_path(old).read_bytes()
                if hashlib.sha256(data).hexdigest() != old.sha256:
                    raise ValueError("local source changed during update")
                path = folder / 'source.docx'
                with path.open('xb') as stream:
                    stream.write(data)
                provenance = {'reused_local_sha256': old.sha256}
            else:
                archive = fetch(row['source_url'])
                with (folder / 'official.zip').open('xb') as stream:
                    stream.write(archive)
                path, provenance = extract_document(archive, spec, folder)
                if path.suffix == '.doc':
                    path = converter(path)
                    provenance['conversion'] = 'legacy_doc_to_docx'
            checksum = sha256_file(path)
            items.append(DocumentSpec(spec, row['latest_version'], path.relative_to(generation).as_posix(),
                                      row['source_url'], checksum))
            sources.append({'spec': spec, 'version': row['latest_version'], 'docx_sha256': checksum, **provenance})
        updated = Manifest(baseline, manifest.release, tuple(items), generation)
        lockfile = generation / 'baseline.lock.json'
        with lockfile.open('x', encoding='utf-8') as stream:
            json.dump(updated.payload(), stream, indent=2)
        with (generation / 'sources.json').open('x', encoding='utf-8') as stream:
            json.dump({'checked_at': catalog.checked_at, 'documents': sources}, stream, indent=2)
        # Reload validates generated paths and manifest invariants before ingestion.
        status = corpus.ingest(load_manifest(lockfile))
        if not status['ingestion_complete']:
            raise ValueError(f"document ingestion failed in {baseline}; active corpus was not changed")
        embedding = {'status': 'not_requested', 'reason': 'explicit_lexical_only'}
        if not lexical_only:
            embedding = {'status': 'complete', 'result': embed(corpus, baseline)}
        new_scopes = list(dict.fromkeys([baseline, *dependencies]))
        new_gaps = reference_inventory(corpus, new_scopes)
        # New documents may reveal further dependencies. Never imply one-hop
        # acquisition established a transitive closure or dated-version coverage.
        available = [doc for scope in new_scopes for doc in corpus.status(scope)['documents']]
        pending = missing_references(new_gaps, available, manifest.release)
        database = corpus.connection.execute('PRAGMA database_list').fetchone()[2]
        result = {'baseline': baseline, 'release': manifest.release, 'manifest': str(lockfile),
                  'manifest_sha256': updated.digest, 'database': str(Path(database).resolve()) if database else None,
                  'checked_at': catalog.checked_at, 'embedding': embedding, 'ingestion_complete': True,
                  'embedding_profile': embedding_profile, 'published': True, 'unchanged': False,
                  'previous_baseline': manifest.baseline, 'missing_references': pending,
                  'unavailable_for_release': [row['spec'] for row in check['documents'] if not row['latest_version']],
                  'global_3gpp_completeness': 'not_assessed', 'document_count': len(items),
                  'active_pointer': str(root / 'current.json')}
        # Publication is last: partial download/parse/embedding leaves old pointer
        # and every prior immutable baseline usable.
        temporary = root / ('.current-' + uuid4().hex + '.json')
        with temporary.open('x', encoding='utf-8') as stream:
            json.dump(result, stream, ensure_ascii=False, indent=2)
            stream.flush()
            os.fsync(stream.fileno())
        os.replace(temporary, root / 'current.json')
        return result
    finally:
        lock.unlink()
