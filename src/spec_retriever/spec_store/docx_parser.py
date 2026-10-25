"""Read source DOCX without destructive cleaning or invented page numbers."""

from __future__ import annotations

from pathlib import Path
import re

from docx import Document
from docx.oxml.ns import qn
from docx.table import Table
from docx.text.paragraph import Paragraph
from docx.enum.style import WD_STYLE_TYPE

from spec_retriever.spec_store.manifest import DocumentSpec
from spec_retriever.spec_store.models import Clause, ParsedDocument

CLAUSE_NUMBER = r"(?:\d+[A-Za-z]?(?:\.\d+[A-Za-z]?)*|[A-Z](?:\.\d+[A-Za-z]?)+)"
NUMBERED = re.compile(r"^(" + CLAUSE_NUMBER + r")\.?\s+(.+)$", re.S)
ANNEX = re.compile(r"^Annex\s+([A-Z])\b[\s:.-]*(.*)$", re.I | re.S)
IDENTITY = re.compile(r"3GPP\s+(?:TS|TR)\s+(\d{2,3}\.\d{3}(?:-\d+)?)\s+(?:V|version\s*)?(\d+\.\d+\.\d+)", re.I)


class _Styles:
    """Resolve immutable document style chains once instead of scanning per block."""

    def __init__(self, document):
        self.styles = {style.style_id: style for style in document.styles if style.type == WD_STYLE_TYPE.PARAGRAPH}
        self.default = document.styles.default(WD_STYLE_TYPE.PARAGRAPH)
        self.chains = {}

    def chain(self, paragraph):
        style_id = paragraph._p.style
        if style_id not in self.chains:
            chain, seen = [], set()
            style = self.styles.get(style_id, self.default)
            while style is not None and style.style_id not in seen:
                seen.add(style.style_id)
                chain.append(style)
                parent = style.element.basedOn_val
                style = self.styles.get(parent) if parent else None
            self.chains[style_id] = tuple(chain)
        return self.chains[style_id]


def _style_chain(paragraph: Paragraph, styles: _Styles | None = None):
    if styles is not None:
        yield from styles.chain(paragraph)
        return
    style = paragraph.style
    seen = set()
    while style is not None and style.style_id not in seen:
        seen.add(style.style_id)
        yield style
        style = style.base_style


def _heading_level(paragraph: Paragraph, styles: _Styles | None = None) -> int | None:
    # In the 3GPP template, TH (table captions) and TT (Contents) can carry
    # outline level 0 too. Outline level alone is not a clause heading.
    if re.match(r"^(?:Table|Figure)\s", paragraph.text, re.I):
        return None
    for style in _style_chain(paragraph, styles):
        match = re.fullmatch(r"heading\s*(\d)", style.name, re.I)
        if match:
            return int(match[1])
    if not (NUMBERED.match(paragraph.text) or ANNEX.match(paragraph.text)):
        return None
    for element in [paragraph._p, *(style.element for style in _style_chain(paragraph, styles))]:
        levels = element.xpath("./w:pPr/w:outlineLvl")
        if levels:
            level = int(levels[0].get(qn("w:val")))
            if level < 9:
                return level + 1
    return None


def _is_toc(paragraph: Paragraph, styles: _Styles | None = None) -> bool:
    return any(re.fullmatch(r"toc\s*\d*", style.name, re.I) for style in _style_chain(paragraph, styles))


class _Numbering:
    """Resolve common Word decimal/letter heading lists, including style links."""

    def __init__(self, document, styles=None):
        self.styles = styles
        try:
            self.root = document.part.numbering_part.element
        except (KeyError, NotImplementedError):
            self.root = None
        self.counts: dict[str, dict[int, int]] = {}
        self.numbers = {} if self.root is None else {
            node.get(qn("w:numId")): node for node in self.root.findall(qn("w:num"))}
        self.abstracts = {} if self.root is None else {
            node.get(qn("w:abstractNumId")): node for node in self.root.findall(qn("w:abstractNum"))}
        self.definitions = {key: {int(node.get(qn("w:ilvl"))): node for node in abstract.findall(qn("w:lvl"))}
                            for key, abstract in self.abstracts.items()}
        self.style_links = []
        for abstract_id, abstract in self.abstracts.items():
            number_id = next((key for key, node in self.numbers.items()
                if node.find(qn("w:abstractNumId")) is not None
                and node.find(qn("w:abstractNumId")).get(qn("w:val")) == abstract_id), None)
            if number_id is not None:
                for definition in abstract.findall(qn("w:lvl")):
                    linked = definition.find(qn("w:pStyle"))
                    if linked is not None:
                        self.style_links.append((linked.get(qn("w:val")), number_id, int(definition.get(qn("w:ilvl")))))
        self.style_link_cache = {}

    def label(self, paragraph: Paragraph, heading_level: int | None) -> str | None:
        if self.root is None:
            return None
        num_id = None
        level = None
        for element in [paragraph._p, *(style.element for style in _style_chain(paragraph, self.styles))]:
            values = element.xpath("./w:pPr/w:numPr/w:numId")
            levels = element.xpath("./w:pPr/w:numPr/w:ilvl")
            if level is None and levels:
                level = int(levels[0].get(qn("w:val")))
            if values:
                num_id = values[0].get(qn("w:val"))
                break
        if num_id == "0":
            return None
        if num_id is None:
            # Some producers attach numbering through abstractNum/lvl/pStyle.
            style_ids = tuple(style.style_id for style in _style_chain(paragraph, self.styles))
            if style_ids not in self.style_link_cache:
                # Preserve the original traversal's last matching style link.
                links = [(number_id, ilvl) for linked, number_id, ilvl in self.style_links if linked in style_ids]
                self.style_link_cache[style_ids] = links[-1] if links else (None, None)
            num_id, level = self.style_link_cache[style_ids]
        if num_id is None:
            return None
        level = level if level is not None else max((heading_level or 1) - 1, 0)
        number = self.numbers.get(num_id)
        if number is None:
            return None
        reference = number.find(qn("w:abstractNumId"))
        if reference is None:
            return None
        abstract_id = reference.get(qn("w:val"))
        abstract = self.abstracts.get(abstract_id)
        if abstract is None:
            return None
        definitions = self.definitions[abstract_id]
        if level not in definitions:
            return None
        counts = self.counts.setdefault(num_id, {})
        definition = definitions[level]
        start = definition.find(qn("w:start"))
        first = int(start.get(qn("w:val"))) if start is not None else 1
        for override in number.findall(qn("w:lvlOverride")):
            if int(override.get(qn("w:ilvl"))) == level:
                override_start = override.find(qn("w:startOverride"))
                if override_start is not None:
                    first = int(override_start.get(qn("w:val")))
        counts[level] = counts[level] + 1 if level in counts else first
        for deeper in list(counts):
            if deeper > level:
                del counts[deeper]
        template = definition.find(qn("w:lvlText"))
        if template is None:
            return None
        label = template.get(qn("w:val"))
        for token in re.findall(r"%(\d+)", label):
            target = int(token) - 1
            if target not in definitions or target not in counts:
                return None  # Never guess a missing parent number.
            fmt = definitions[target].find(qn("w:numFmt"))
            fmt = fmt.get(qn("w:val")) if fmt is not None else "decimal"
            value = counts[target]
            if fmt == "decimal":
                rendered = str(value)
            elif fmt in {"upperLetter", "lowerLetter"} and 1 <= value <= 26:
                rendered = chr((65 if fmt == "upperLetter" else 97) + value - 1)
            else:
                return None
            label = label.replace("%" + token, rendered)
        return label.strip(" .\t")


def _table_block(table: Table, location: dict) -> dict:
    rows = []
    # Use physical XML cells to avoid falsely duplicating merged-cell values.
    for row in table.rows:
        cells = []
        for cell in row._tr.tc_lst:
            text = []
            for child in cell:
                if child.tag == qn("w:p"):
                    text.append(Paragraph(child, table).text)
                elif child.tag == qn("w:tbl"):
                    text.append(_table_block(Table(child, table), location)["text"])
            props = cell.find(qn("w:tcPr"))
            span = props.find(qn("w:gridSpan")) if props is not None else None
            merge = props.find(qn("w:vMerge")) if props is not None else None
            cells.append({"text": "\n".join(text),
                          "column_span": int(span.get(qn("w:val"))) if span is not None else 1,
                          "vertical_merge": (merge.get(qn("w:val")) or "continue") if merge is not None else None})
        rows.append(cells)
    return {"kind": "table", "rows": rows, "location": location,
            "text": "\n".join("\t".join(cell["text"] for cell in row) for row in rows)}


def parse_docx(path: Path, expected: DocumentSpec) -> ParsedDocument:
    document = Document(path)
    styles = _Styles(document)
    warnings = []
    body = document.element.body
    unsupported = {
        "images_not_interpreted": ".//w:drawing | .//w:pict",
        "equations_not_interpreted": ".//m:oMath",
        "revisions_not_interpreted": ".//w:ins | .//w:del",
        "text_boxes_not_interpreted": ".//w:txbxContent",
        "content_controls_not_interpreted": ".//w:sdt",
        "footnotes_not_interpreted": ".//w:footnoteReference | .//w:endnoteReference",
        "embedded_objects_not_interpreted": ".//w:object | .//w:altChunk",
    }
    for code, xpath in unsupported.items():
        count = len(body.xpath(xpath))
        if count:
            warnings.append({"code": code, "count": count})
    # Inspect the cover, not references in the specification body: those may
    # correctly name other specifications and older Releases.
    cover = []
    for element in body:
        if element.tag == qn("w:tbl"):
            cover.append(" ".join(element.xpath(".//w:t/text()")))
            continue
        if element.tag != qn("w:p"):
            continue
        paragraph = Paragraph(element, document)
        if _heading_level(paragraph, styles) is not None:
            break
        cover.append(paragraph.text)
        if len(cover) == 100:
            break
    cover_identities = set(IDENTITY.findall("\n".join(cover)))
    header_footer = []
    for section in document.sections:
        header_footer.extend(paragraph.text for paragraph in section.header.paragraphs)
        header_footer.extend(paragraph.text for paragraph in section.footer.paragraphs)
    header_identities = set(IDENTITY.findall("\n".join(header_footer)))
    # STYLEREF/revision fields can concatenate current and previous header
    # versions. The explicit cover is authoritative when it identifies a spec.
    identities = cover_identities or header_identities
    if cover_identities and header_identities - cover_identities:
        warnings.append({"code": "header_identity_differs_from_cover", "message": "cover identity used; header fields require review"})
    if identities:
        if identities != {(expected.spec, expected.version)}:
            raise ValueError(f"document identity {sorted(identities)} differs from manifest {expected.spec} {expected.version}")
    else:
        warnings.append({"code": "document_identity_not_detected", "message": "identity relies on manifest metadata and checksum"})

    current = Clause("preamble", "Preamble", ["Preamble"], None, 0)
    clauses = [current]
    seen = {"preamble"}
    stack: list[tuple[int, Clause]] = []
    numbering = _Numbering(document, styles)
    for body_index, element in enumerate(body):
        location = {"part": "word/document.xml", "body_index": body_index}
        if element.tag == qn("w:tbl"):
            current.blocks.append(_table_block(Table(element, document), location))
            continue
        if element.tag != qn("w:p"):
            continue
        paragraph = Paragraph(element, document)
        text = paragraph.text.strip()
        level = _heading_level(paragraph, styles)
        label = numbering.label(paragraph, level)
        annex = ANNEX.match(text) if not _is_toc(paragraph, styles) else None
        if annex and level is None and not re.match(r"^\s*\((?:normative|informative)\)\s*:", annex.group(2), re.I):
            annex = None  # Body citations and change-history lines are not headings.
        if _is_toc(paragraph, styles):
            level = None
        if text and (level is not None or annex):
            match = NUMBERED.match(text)
            annex_type = None
            if annex:
                clause_id, title = annex.group(1).upper(), text
                level = 1
            elif match:
                clause_id, title = match.groups()
            elif label and re.fullmatch(r"\d+(?:\.\d+)*|[A-Z](?:\.\d+)*", label):
                clause_id, title = label, text
            else:
                # Unnumbered front matter is retained without fabricated IDs.
                clause_id, title = f"heading-{body_index}", text
                warnings.append({"code": "unresolved_heading_number", "location": location, "title": text})
            if clause_id in seen:
                # Some published DOCX files contain the same heading twice in
                # adjacent paragraphs. Keep one clause when the first heading
                # has no body yet; any other collision may hide real content.
                if (current.clause_id == clause_id and current.title == title
                        and len(current.blocks) == 1
                        and current.blocks[0]["kind"] == "heading"
                        and current.blocks[0]["text"] == text
                        and current.blocks[0]["location"]["body_index"] == body_index - 1):
                    warnings.append({"code": "duplicate_identical_heading",
                                     "clause_id": clause_id,
                                     "first_location": current.blocks[0]["location"],
                                     "duplicate_location": location})
                    continue
                printed_id = clause_id
                occurrence = 2
                while f"{printed_id}~{occurrence}" in seen:
                    occurrence += 1
                clause_id = f"{printed_id}~{occurrence}"
                warnings.append({"code": "duplicate_printed_clause_id",
                                 "printed_clause_id": printed_id,
                                 "resolved_clause_id": clause_id,
                                 "location": location,
                                 "message": "distinct source headings share a printed number; both are retained"})
            seen.add(clause_id)
            level = level or clause_id.count(".") + 1
            while stack and stack[-1][0] >= level:
                stack.pop()
            if re.fullmatch(r"[A-Z](?:\.\d+)*", clause_id):
                tag = re.search(r"\((normative|informative)\)", text, re.I)
                annex_type = tag.group(1).lower() if tag else (stack[-1][1].annex_type if stack else "unspecified")
            parent = stack[-1][1] if stack else None
            title_path = [*(parent.title_path if parent else []), f"{clause_id} {title}"]
            current = Clause(clause_id, title, title_path, parent.clause_id if parent else None,
                             len(clauses), annex_type)
            current.blocks.append({"kind": "heading", "text": text, "location": location})
            clauses.append(current)
            stack.append((level, current))
        elif text:
            current.blocks.append({"kind": "paragraph", "text": text, "location": location})
            if current.annex_type == "unspecified":
                tag = re.match(r"^\((normative|informative)\)", text, re.I)
                if tag:
                    current.annex_type = tag.group(1).lower()

    if not any(clause.blocks for clause in clauses):
        raise ValueError("DOCX contains no supported textual blocks")
    if len(clauses) == 1:
        warnings.append({"code": "no_clause_headings_detected"})
    return ParsedDocument(clauses, warnings)
